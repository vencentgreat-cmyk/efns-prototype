"""Synthetic associated-view evidence; no real company export is a fixture."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import datetime, date
from io import BytesIO
from types import SimpleNamespace
from urllib.parse import urlencode
from uuid import UUID

from openpyxl import Workbook
import pandas as pd
import pytest

from app.security import Role, User, AuthorizationError
from app.services.authorized_repository import AuthorizedRepository
from data.flock_transaction_supplement import parse_associated_view, SupplementError
from data.repositories.base import RepositoryError, ConcurrencyError
from data.repositories.mock import MockRepository


def ident(number): return str(UUID(int=number)).upper()
FLOCK, ACCOUNT, FACILITY, ANCHOR = (ident(n) for n in (1,2,3,1000))
BATCH='DEV_MIGRATION_SYNTHETIC_FT'
ACTOR={'user_id':ident(700),'email':'synthetic.operator@nsegg.ca'}
BINDING=dict(original_batch_id=BATCH,flock_id=FLOCK,anchor_transaction_id=ANCHOR,
             basis='Synthetic associated-view export and same-GUID Delivery anchor reviewed.')


def workbook(mutate=None):
    book=Workbook(); sheet=book.active; sheet.title='Synthetic associated view'
    headers=['Source GUID','Checksum','Modified','Permit','Source date','Birds','Type','Count origin','Notes','Created']
    sheet.append(headers)
    for i in range(13):
        sheet.append([ident(1000+i),'synthetic-checksum',datetime(2025,4,1),'P-SYN',
            datetime(2025,3,1+i,1),10+i,'Delivery' if i==0 else 'Count',
            None if i==0 else ('Board' if i==1 else 'Producer'),None,datetime(2025,3,20)])
    hidden=book.create_sheet('Metadata');hidden.sheet_state='veryHidden'
    fields=dict(zip(('new_flocktransactionid','checksumLogicalName','modifiedon','new_flockidid',
        'new_date','new_birdcount','new_flocktransactiontype','efc_flockcounttype','new_comments','createdon'),headers))
    hidden['A1']='new_flocktransaction:synthetic-checksum:'+urlencode(fields)
    if mutate: mutate(sheet,hidden)
    out=BytesIO();book.save(out);return out.getvalue()


@pytest.fixture
def repo():
    r=MockRepository(seed=2)
    r._accounts=pd.DataFrame([dict(ACCOUNT_ID=ACCOUNT,ORGANIZATION_NAME='Synthetic Account')])
    r._facilities=pd.DataFrame([dict(FACILITY_ID=FACILITY,ACCOUNT_ID=ACCOUNT,FACILITY_NAME='Synthetic Facility')])
    r._flocks=pd.DataFrame([dict(FLOCK_ID=FLOCK,ACCOUNT_ID=ACCOUNT,FACILITY_ID=FACILITY,
        FLOCK_NUMBER='F-SYN',PERMIT_NUMBER='P-SYN',BIRD_COUNT=10,STATUS='Planned',
        PLACEMENT_DATE=None,DISPOSAL_DATE=None,UPDATED_AT=datetime(2025,1,1))])
    anchor=dict(FLOCK_TRANSACTION_ID=ANCHOR,FLOCK_ID=FLOCK,TRANSACTION_TYPE='Delivery',
                QUANTITY=10,TRANSACTION_DATE=date(2025,3,1),NOTES=None)
    r._flock_transactions=pd.DataFrame([anchor])
    r._migration_batches=[dict(MIGRATION_BATCH_ID=BATCH,STATUS='COMMITTED')]
    r._migration_raw_rows=[dict(MIGRATION_BATCH_ID=BATCH,MIGRATION_RAW_ROW_ID='synthetic-anchor',
        SOURCE_ENTITY='FLOCK_TRANSACTION',SOURCE_ID=ANCHOR,VALIDATION_STATUS='READY',NORMALIZED_DATA=anchor)]
    r._migration_id_map=[dict(MIGRATION_BATCH_ID=BATCH,MIGRATION_RAW_ROW_ID='synthetic-anchor',
        SOURCE_ENTITY='FLOCK_TRANSACTION',SOURCE_ID=ANCHOR,TARGET_ID=ANCHOR,STATUS='COMMITTED')]
    return r


def preview(repo,content=None,binding=None):
    return repo.preview_flock_transaction_supplement('synthetic.xlsx',content or workbook(),binding or BINDING)


def commit(repo,content=None,review=None,binding=None):
    content=content or workbook();binding=binding or BINDING
    review=review or preview(repo,content,binding)
    return repo.commit_flock_transaction_supplement('synthetic.xlsx',content,binding,
        review['review_token'],confirmed=True,actor=ACTOR)


def test_first_commit_and_retry_keep_all_evidence_and_flock_unchanged(repo):
    content=workbook();before=deepcopy((repo._migration_batches,repo._migration_raw_rows,repo._migration_id_map))
    flock=repo._flocks.copy(deep=True);other=repo._production.copy(deep=True)
    report=preview(repo,content)
    assert report['counts']==dict(NEW=12,EXISTING_MATCH=1,CONFLICT=0,INVALID=0)
    assert not repo.get_source_files().shape[0] and len(repo._flock_transactions)==1
    result=commit(repo,content,report)
    assert (result['created_count'],result['existing_count'],result['updated_count'])==(12,1,0)
    raw=repo.get_raw_rows(result['import_id'])
    assert len(raw)==13 and set(raw.VALIDATION_STATUS)=={'NEW','EXISTING_MATCH'}
    assert raw.iloc[0]['RAW_DATA']['cells']['E2']['value']==datetime(2025,3,1,1)
    assert 'xml' in raw.iloc[0]['RAW_DATA']['cells']['E2']
    assert repo.get_source_file_bytes(raw.iloc[0]['SOURCE_FILE_ID'])==content
    assert len(repo._supplement_audit)==1 and repo._supplement_audit[0]['ACTOR_EMAIL']==ACTOR['email']
    assert repo._flock_transactions.COUNT_SOURCE_LABEL.value_counts().to_dict()=={'Producer':11,'Board':1}
    after=preview(repo,content)
    assert after['counts']==dict(NEW=0,EXISTING_MATCH=13,CONFLICT=0,INVALID=0)
    again=commit(repo,content,after)
    assert again['created_count']==0 and again['existing_count']==13 and again['already_committed']
    assert len(repo._flock_transactions)==13 and len(repo._supplement_audit)==1
    pd.testing.assert_frame_equal(flock,repo._flocks);pd.testing.assert_frame_equal(other,repo._production)
    assert before==(repo._migration_batches,repo._migration_raw_rows,repo._migration_id_map)


@pytest.mark.parametrize('field,value',[('QUANTITY',99),('FLOCK_ID',ident(900)),
    ('COUNT_SOURCE_LABEL','Different'),('TRANSACTION_DATE',date(2025,2,1)),('NOTES','Different')])
def test_same_guid_conflict_blocks_entire_batch(repo,field,value):
    parsed=parse_associated_view('synthetic.xlsx',workbook())
    row=dict(parsed.rows[1]['record'],FLOCK_ID=FLOCK);row[field]=value
    repo._flock_transactions=pd.concat([repo._flock_transactions,pd.DataFrame([row])],ignore_index=True)
    report=preview(repo)
    assert report['counts']['CONFLICT']==1 and not report['can_commit']
    with pytest.raises(RepositoryError): commit(repo,review=report)
    assert len(repo._flock_transactions)==2 and repo.get_source_files().empty


@pytest.mark.parametrize('column,value',[(1,None),(1,'not-a-guid'),(5,'bad-date'),(6,'bad-number'),
    (6,-1),(6,1.234),(7,'Unsupported'),(8,None),(6,'=1+1')])
def test_invalid_source_rows_visible_and_block_commit(repo,column,value):
    content=workbook(lambda s,_:setattr(s.cell(3,column),'value',value))
    report=preview(repo,content)
    assert report['counts']['INVALID']>=1 and len(report['rows'])==13 and not report['can_commit']
    with pytest.raises(RepositoryError): commit(repo,content,report)
    assert len(repo._flock_transactions)==1


def test_duplicate_guid_is_invalid_not_deduplicated(repo):
    content=workbook(lambda s,_:setattr(s['A4'],'value',s['A3'].value))
    report=preview(repo,content)
    assert report['counts']['INVALID']==2 and len(report['rows'])==13


def test_anchor_mismatch_is_visible_conflict_not_an_update(repo):
    repo._flock_transactions.loc[0,'QUANTITY']=11
    report=preview(repo,workbook())
    assert report['rows'][0]['classification']=='CONFLICT' and not report['can_commit']
    with pytest.raises(RepositoryError): commit(repo,workbook(),report)
    assert repo._flock_transactions.iloc[0]['QUANTITY']==11


def test_unknown_source_label_is_preserved_not_recoded(repo):
    content=workbook(lambda s,_:setattr(s['H3'],'value','Unconfirmed source label'))
    assert preview(repo,content)['can_commit']
    commit(repo,content)
    assert repo.get_flock_transactions().iloc[1]['COUNT_SOURCE_LABEL']=='Unconfirmed source label'


@pytest.mark.parametrize('change',['no_binding','no_basis','missing_target','target_changed','anchor_changed','wrong_lineage','lookup_conflict'])
def test_binding_and_review_protection(repo,change):
    content=workbook();binding=dict(BINDING);report=preview(repo,content)
    if change=='no_binding': binding.pop('flock_id')
    if change=='no_basis': binding['basis']=''
    if change=='missing_target': repo._flocks=repo._flocks.iloc[:0]
    if change=='target_changed': repo._flocks.loc[0,'STATUS']='Active'
    if change=='anchor_changed': repo._flock_transactions.loc[0,'QUANTITY']=11
    if change=='wrong_lineage': repo._migration_id_map[0]['MIGRATION_BATCH_ID']='another-batch'
    if change=='lookup_conflict': content=workbook(lambda s,_:setattr(s['D3'],'value','different-lookup'))
    with pytest.raises((SupplementError,RepositoryError)):
        commit(repo,content,report,binding)
    assert len(repo._flock_transactions)==1


@pytest.mark.parametrize('failure',['insert','raw','audit'])
def test_partial_failure_rolls_back_core_tracking_audit_but_retains_archive(repo,monkeypatch,failure):
    if failure=='insert':
        original=repo._insert_supplement_transaction;counter=[]
        def broken(tx,row):
            original(tx,row);counter.append(1)
            if len(counter)==4: raise RuntimeError('synthetic write failure')
        monkeypatch.setattr(repo,'_insert_supplement_transaction',broken)
    elif failure=='raw': monkeypatch.setattr(repo,'insert_raw_rows',lambda *_: (_ for _ in ()).throw(RuntimeError('raw failure')))
    else: monkeypatch.setattr(repo,'_supplement_audit_event',lambda *_: (_ for _ in ()).throw(RuntimeError('audit failure')))
    before=len(repo._import_batches)
    with pytest.raises(RuntimeError): commit(repo)
    assert len(repo._flock_transactions)==1 and len(repo._import_batches)==before and not repo._raw_rows
    assert not getattr(repo,'_supplement_audit',[])
    assert len(repo.get_source_files())==1 and repo.get_source_files().iloc[0]['PARSE_STATUS']=='ARCHIVED'


def test_concurrent_same_file_creates_only_once(repo):
    content=workbook();report=preview(repo,content)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(lambda _:commit(repo,content,report),range(2)))
    assert sorted(row['created_count'] for row in results)==[0,12]
    assert len(repo._flock_transactions)==13 and len(repo._supplement_audit)==1


@pytest.mark.parametrize('role',[Role.REPORTING_VIEWER,Role.DATA_EDITOR,Role.DEVELOPER])
def test_only_admin_authorized_for_preview_or_commit(repo,role):
    user=User(ident(700),'synthetic.operator@nsegg.ca','Synthetic',role,True,False,'','',None)
    secured=AuthorizedRepository(repo,SimpleNamespace(),user)
    with pytest.raises(RepositoryError): secured.preview_flock_transaction_supplement('synthetic.xlsx',workbook(),BINDING)
    with pytest.raises(RepositoryError): secured.commit_flock_transaction_supplement('synthetic.xlsx',workbook(),BINDING,'token',confirmed=True)


def test_admin_actor_is_injected_not_frontend_and_confirmation_required(repo):
    user=User(ACTOR['user_id'],ACTOR['email'],'Synthetic',Role.ADMIN,True,False,'','',None)
    secured=AuthorizedRepository(repo,SimpleNamespace(),user)
    content=workbook();report=secured.preview_flock_transaction_supplement('synthetic.xlsx',content,BINDING)
    with pytest.raises(RepositoryError):
        secured.commit_flock_transaction_supplement('synthetic.xlsx',content,BINDING,report['review_token'])
    secured.commit_flock_transaction_supplement('synthetic.xlsx',content,BINDING,report['review_token'],confirmed=True,
        actor={'user_id':'forged','email':'forged@example.org'})
    assert repo._supplement_audit[0]['ACTOR_EMAIL']==ACTOR['email']


def test_report_preview_csv_xlsx_preserve_source_labels_and_old_null(repo):
    from app.services.report_center import build_custom_report,report_csv,report_xlsx
    user=User(ACTOR['user_id'],ACTOR['email'],'Synthetic',Role.ADMIN,True,False,'','',None)
    commit(repo,workbook())
    frame=build_custom_report(repo,user=user,dataset='Flock Transactions',
        selected_fields=['FLOCK_TRANSACTION_ID','COUNT_SOURCE_LABEL','QUANTITY'])
    assert len(frame)==13 and frame['Count Source'].isna().sum()==1
    assert frame['Count Source'].value_counts().to_dict()=={'Producer':11,'Board':1}
    csv_frame=pd.read_csv(BytesIO(report_csv(user,frame)))
    excel_frame=pd.read_excel(BytesIO(report_xlsx(user,frame)))
    assert csv_frame['Count Source'].fillna('').tolist()==frame['Count Source'].fillna('').tolist()
    assert excel_frame['Count Source'].fillna('').tolist()==frame['Count Source'].fillna('').tolist()


def test_changed_source_and_binding_require_new_review(repo):
    content=workbook();report=preview(repo,content)
    modified=workbook(lambda s,_:setattr(s['F3'],'value',999))
    with pytest.raises(ConcurrencyError): commit(repo,modified,report)
    changed=dict(BINDING,basis='A different review decision')
    with pytest.raises(ConcurrencyError): commit(repo,content,report,changed)
    assert not repo.get_source_files().shape[0]


def test_archive_tampering_rejected_before_core_write(repo,monkeypatch):
    monkeypatch.setattr(repo,'get_source_file_bytes',lambda _: b'tampered')
    with pytest.raises(RepositoryError,match='differs'): commit(repo,workbook())
    assert len(repo._flock_transactions)==1 and not repo._raw_rows


def test_failed_write_can_retry_same_archive_without_duplicates(repo,monkeypatch):
    content=workbook();review=preview(repo,content)
    original=repo._insert_supplement_transaction
    counter=[]
    def fail_once(tx,record):
        original(tx,record);counter.append(1)
        if len(counter)==2: raise RuntimeError('synthetic failed attempt')
    monkeypatch.setattr(repo,'_insert_supplement_transaction',fail_once)
    with pytest.raises(RuntimeError): commit(repo,content,review)
    monkeypatch.setattr(repo,'_insert_supplement_transaction',original)
    result=commit(repo,content,review)
    assert result['created_count']==12 and len(repo._flock_transactions)==13
    assert len(repo.get_source_files())==1 and len(repo._raw_rows)==13


def test_target_change_between_archive_and_locked_recheck_is_blocked(repo,monkeypatch):
    original=repo.archive_source_file
    def archive_then_external_change(*args,**kwargs):
        result=original(*args,**kwargs)
        repo._flocks.loc[0,'STATUS']='Active'
        return result
    monkeypatch.setattr(repo,'archive_source_file',archive_then_external_change)
    with pytest.raises(ConcurrencyError): commit(repo,workbook())
    assert len(repo._flock_transactions)==1 and not repo._raw_rows
