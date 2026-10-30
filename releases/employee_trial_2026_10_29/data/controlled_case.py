"""Restricted October-29 source case. Pure parsing; never connects or writes."""
from io import BytesIO
from pathlib import PurePosixPath
from collections import Counter
from urllib.parse import parse_qsl,urlparse
from zipfile import ZipFile
import hashlib,json,uuid
from xml.etree import ElementTree as ET
import pandas as pd
from openpyxl import load_workbook
from data.migration import load_new_package_mapping,normalize_migration_row
from data.source_evidence import evidence_json
from data.repositories.base import RepositoryError

VERSION='EFNS_OCT29_INCREMENTAL_1'
ORG='eggfarmersofns.crm.dynamics.com'
MAX_BIND=1024*1024

def digest(value):return hashlib.sha256(evidence_json(value).encode()).hexdigest()
def bytehash(value):return hashlib.sha256(bytes(value)).hexdigest()
def guid(value):
    try:return str(uuid.UUID(str(value))).upper()
    except Exception as exc:raise RepositoryError('A real source GUID is required.') from exc
def optional_guid(value):return None if value is None else guid(value)
def unique(rows,key):
    result={key(r):r for r in rows}
    if len(result)!=len(rows):raise RepositoryError('Duplicate source or target identity; no count-only verification.')
    return result
def batches(rows,max_rows=500,max_bytes=MAX_BIND):
    result=[];current=[];size=2
    for r in rows:
        length=len(evidence_json(r).encode())
        if length+2>max_bytes:raise RepositoryError('A single evidence row exceeds the safe bind size.')
        if current and (len(current)==max_rows or size+length+2>max_bytes):result.append(current);current=[];size=2
        size+=length+(2 if current else 0);current.append(r)
    if current:result.append(current)
    return result

def json_document(content):
    if len(content)>16*1024*1024:raise RepositoryError('Source JSON exceeds this restricted case limit.')
    doc=json.loads(bytes(content).decode('utf-8-sig'))
    if urlparse(doc.get('@odata.context','')).hostname!=ORG:raise RepositoryError('Source API organization mismatch.')
    rows=doc.get('value')
    if not isinstance(rows,list):raise RepositoryError('A source API value array is required.')
    return doc

def api_rows(content,zipped=False):
    if not zipped:
        doc=json_document(content)
        if doc.get('@odata.nextLink'):raise RepositoryError('Unfinished single-document source API export.')
        return doc['value']
    rows=[]
    with ZipFile(BytesIO(content)) as z:
        members=sorted(z.infolist(),key=lambda m:m.filename)
        if len(members)!=4 or sum(m.file_size for m in members)>16*1024*1024:raise RepositoryError('Only the four reviewed JSON pages are supported.')
        for i,m in enumerate(members):
            if m.is_dir() or PurePosixPath(m.filename).name!=m.filename or not m.filename.endswith('.json'):raise RepositoryError('Unsafe source archive member.')
            doc=json_document(z.read(m));link=doc.get('@odata.nextLink')
            if bool(link)!=(i<3) or link and urlparse(link).hostname!=ORG:raise RepositoryError('Incomplete or wrong-organization API pagination.')
            rows.extend(doc['value'])
    return rows

def workbook_rows(content,entity):
    if len(content)>20*1024*1024:raise RepositoryError('Workbook exceeds the restricted source limit.')
    with ZipFile(BytesIO(content)) as z:
        if sum(m.file_size for m in z.infolist())>128*1024*1024:raise RepositoryError('Workbook expansion exceeds the case limit.')
        for name in z.namelist():
            if name.endswith('.rels') and any(n.get('TargetMode')=='External' for n in ET.fromstring(z.read(name))):raise RepositoryError('External workbook relationships are not executed or accepted.')
    wb=load_workbook(BytesIO(content),data_only=False,keep_links=False)
    visible=[s for s in wb if s.sheet_state=='visible'];hidden=[s for s in wb if s.sheet_state!='visible']
    if len(visible)!=1 or len(hidden)!=1:raise RepositoryError('Expected one source sheet and one metadata sheet.')
    descriptor=str(hidden[0]['A1'].value or '');parts=descriptor.split(':',2)
    logical_entity='new_flocks' if entity=='FLOCK' else 'new_flocktransaction'
    if len(parts)!=3 or parts[0]!=logical_entity:raise RepositoryError('Wrong source entity metadata.')
    pairs=parse_qsl(parts[2]);headers=[c.value for c in visible[0][1]]
    if len(headers)!=len(set(headers)):raise RepositoryError('Duplicate source headings.')
    inverse={label:key for key,label in pairs};pk='new_flocksid' if entity=='FLOCK' else 'new_flocktransactionid';rows=[]
    for n,cells in enumerate(visible[0].iter_rows(min_row=2),2):
        if not any(c.value is not None for c in cells):continue
        if any(c.data_type=='f' for c in cells):raise RepositoryError('Source formulas are not evaluated or accepted.')
        logical={inverse[h]:c.value for h,c in zip(headers,cells) if h in inverse}
        if pk not in logical:raise RepositoryError('Missing source primary GUID column.')
        rows.append({'guid':guid(logical[pk]),'sheet':visible[0].title,'row':n,'logical':logical,'cells':[{'column':c.coordinate,'value':c.value,'type':c.data_type,'format':c.number_format} for c in cells],'checksum':next((c.value for h,c in zip(headers,cells) if h=='(Do Not Modify) Row Checksum'),None),'system_modifiedon':next((c.value for h,c in zip(headers,cells) if h=='(Do Not Modify) Modified On'),None),'metadata':descriptor})
    return unique(rows,lambda r:r['guid'])

def build_case(blobs,policy):
    """Policy contains reviewed exact source IDs, never arbitrary CORE assignments."""
    if set(policy)!={'original_batch','files','phase_a_ids','phase_b_ids','hold_ids','excluded_ids'}:raise RepositoryError('Unexpected case policy fields.')
    a=[guid(x) for x in policy['phase_a_ids']];b=[guid(x) for x in policy['phase_b_ids']];holds=[guid(x) for x in policy['hold_ids']];excluded=[guid(x) for x in policy['excluded_ids']]
    if not isinstance(policy['hold_ids'],dict) or any(not isinstance(v,str) or not v.strip() for v in policy['hold_ids'].values()):raise RepositoryError('Each reviewed HOLD must retain its explicit reason.')
    if not a or not b or len(a)>9 or len(b)>1991 or len(holds)>43 or len(set(a+b+holds+excluded))!=len(a+b+holds+excluded):raise RepositoryError('Scope overlaps or exceeds the restricted reviewed case.')
    specs=policy['files'];kinds={f['kind']:f for f in specs}
    if len(kinds)!=len(specs) or set(kinds)!={'FLOCK_XLSX','TRANSACTION_XLSX','FLOCK_API','TRANSACTION_API','STATUS_API','NULL_API'}:raise RepositoryError('Exactly the six reviewed source kinds are required.')
    for f in specs:
        if set(f)!={'kind','name','sha256','size','origin'} or PurePosixPath(f['name']).name!=f['name']:raise RepositoryError('Unapproved source specification.')
        suffixes={k:({'.xlsx'} if k.endswith('XLSX') else {'.zip'} if k=='TRANSACTION_API' else {'.json','.txt'}) for k in kinds}
        if PurePosixPath(f['name']).suffix.lower() not in suffixes[f['kind']]:raise RepositoryError('Source file extension does not match its restricted kind.')
        content=blobs[f['name']]
        if len(content)!=f['size'] or bytehash(content)!=f['sha256']:raise RepositoryError('Original source bytes changed; do not rebase.')
    def content(kind):return blobs[kinds[kind]['name']]
    xf=workbook_rows(content('FLOCK_XLSX'),'FLOCK');xt=workbook_rows(content('TRANSACTION_XLSX'),'FLOCK_TRANSACTION')
    af=unique(api_rows(content('FLOCK_API')),lambda r:guid(r['new_flocksid']));at=unique(api_rows(content('TRANSACTION_API'),True),lambda r:guid(r['new_flocktransactionid']))
    status=unique(api_rows(content('STATUS_API')),lambda r:guid(r['new_flocktransactionid']))
    if set(status)!=set(excluded) or any(r.get('statecode')!=1 or r.get('statuscode')!=2 for r in status.values()):raise RepositoryError('The excluded 56/status scope changed.')
    nulls=unique(api_rows(content('NULL_API')),lambda r:guid(r['new_flocktransactionid']))
    if not set(nulls)<=set(holds) or any(r.get('new_date') is not None for r in nulls.values()):raise RepositoryError('Confirmed source-NULL holds changed; new approval is needed.')
    for id,r in {**status,**nulls}.items():
        if id not in at or r['@odata.etag']!=at[id]['@odata.etag'] or r['modifiedon']!=at[id]['modifiedon']:raise RepositoryError('API evidence versions disagree.')
    mapping=load_new_package_mapping();records=[]
    for entity,ids,xlsx,api,kind in [('FLOCK',a,xf,af,'FLOCK_XLSX'),('FLOCK_TRANSACTION',b+holds,xt,at,'TRANSACTION_XLSX')]:
        for id in sorted(ids):
            x=xlsx[id];r=api[id];clock=pd.Timestamp(r['modifiedon'])
            if clock.tzinfo is None or clock.tz_convert('UTC').tz_localize(None)!=pd.Timestamp(x['system_modifiedon']):raise RepositoryError('Source version/clock representation mismatch; do not merge epochs.')
            if not r.get('@odata.etag'):raise RepositoryError('Source etag is required.')
            values=dict(x['logical'])
            for field in mapping['entities'][entity]['field_map'].values():
                if field.endswith('_displayname') and field[:-12] in values:values[field]=values[field[:-12]]
                elif field.endswith('name') and field[:-4] in values:values[field]=values[field[:-4]]
            if entity=='FLOCK':values.update(new_flocksid=id,new_account=optional_guid(r['_new_account_value']),new_facilityidid=optional_guid(r['_new_facilityidid_value']))
            else:values.update(new_flocktransactionid=id,new_flockidid=optional_guid(r['_new_flockidid_value']))
            core,errors=normalize_migration_row(entity,values,x['row'],mapping)
            for field in mapping['entities'][entity]['required']:
                if core.get(field) is None:errors.append(field+' is missing')
            phase='HOLD' if id in holds else 'A' if entity=='FLOCK' else 'B'
            if errors and phase!='HOLD':raise RepositoryError('A reviewed eligible source has invalid business fields.')
            if phase=='HOLD':errors.append(policy['hold_ids'][id])
            records.append({'entity':entity,'guid':id,'phase':phase,'core':core,'issues':errors,'source':{'xlsx_file':kinds[kind]['name'],'xlsx_sha256':kinds[kind]['sha256'],'xlsx':x,'api_file':kinds['FLOCK_API' if entity=='FLOCK' else 'TRANSACTION_API']['name'],'api_sha256':kinds['FLOCK_API' if entity=='FLOCK' else 'TRANSACTION_API']['sha256'],'api':r,'optional_lookup_state':'UNVERIFIED_LABELS_RETAINED_NO_GUID_INFERENCE'}})
    body={'version':VERSION,'policy':policy,'mapping_hash':digest(mapping),'records':records}
    return {**body,'case_hash':digest(body),'operation_id':str(uuid.uuid5(uuid.NAMESPACE_URL,VERSION+':'+digest(body)))}

def validate_case(plan):
    body={k:plan[k] for k in ('version','policy','mapping_hash','records')}
    if plan['version']!=VERSION or digest(body)!=plan['case_hash'] or digest(load_new_package_mapping())!=plan['mapping_hash'] or plan['operation_id']!=str(uuid.uuid5(uuid.NAMESPACE_URL,VERSION+':'+plan['case_hash'])):raise RepositoryError('Case, operation identity or mapping seal changed.')
    expected={(e,guid(id),phase) for e,ids,phase in [('FLOCK',plan['policy']['phase_a_ids'],'A'),('FLOCK_TRANSACTION',plan['policy']['phase_b_ids'],'B'),('FLOCK_TRANSACTION',plan['policy']['hold_ids'],'HOLD')] for id in ids}
    actual={(r['entity'],r['guid'],r['phase']) for r in plan['records']}
    if actual!=expected or len(actual)!=len(plan['records']):raise RepositoryError('Case record scope changed.')
    return plan
