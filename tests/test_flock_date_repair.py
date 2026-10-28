"""Synthetic date-contract correction tests; never connect to Snowflake."""
from copy import deepcopy
from datetime import date,datetime
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from types import SimpleNamespace
from uuid import UUID
import json
import pandas as pd
import pytest
from openpyxl import load_workbook

from data.migration import load_mapping, load_mapping_for_version, V3_MAPPING_PATH, normalize_migration_row
from data.operational_flock_quota_import import _mapping, _aliases
from data.repositories.mock import MockRepository
from data.repositories.base import RepositoryError, ConcurrencyError
from data.repositories.flock_date_repair import KIND, UNDO
from app.security import Role
from app.services.authorized_repository import AuthorizedRepository
from app.services.report_center import build_custom_report,report_csv,report_xlsx
from tests.test_report_center import user

BATCH='DEV_MIGRATION_SYNTHETIC_DATES'
ACTOR={'user_id':str(UUID(int=900)),'email':'synthetic.reviewer@nsegg.ca'}
BASIS='Synthetic verified source metadata and page mapping; no lifecycle change.'


@pytest.fixture
def repo():
    r=MockRepository(seed=3)
    original=load_mapping(V3_MAPPING_PATH)
    r._migration_batches=[dict(MIGRATION_BATCH_ID=BATCH,STATUS='COMMITTED',SCHEMA_VERSION=original['schema_version'])]
    r._migration_raw_rows=[];r._migration_id_map=[]
    rows=[]
    for n,(actual,estimated) in enumerate([('2025-05-01','2025-06-01'),(None,'2025-08-01'),(None,None),('2025-09-01','2025-09-01')],1):
        g=str(UUID(int=n)).upper()
        raw=dict(NEW_FLOCKSID=g,NEW_DISPOSALDATE=actual,NEW_DISPOSALDATE2=estimated)
        normalized,_=normalize_migration_row('FLOCK',raw,n+1,original)
        rows.append(dict(normalized,FLOCK_ID=g,FLOCK_NUMBER=f'SYN-{n}',PERMIT_NUMBER=f'P-{n}',
            ACCOUNT_ID='synthetic-account',FACILITY_ID='synthetic-facility',STATUS='Active',BIRD_COUNT=100,
            PLACEMENT_DATE=date(2025,1,1),HATCH_DATE=date(2024,10,1),
            CREATED_AT=datetime(2025,10,1),UPDATED_AT=datetime(2025,10,1)))
        r._migration_raw_rows.append(dict(MIGRATION_BATCH_ID=BATCH,MIGRATION_RAW_ROW_ID=f'raw-{n}',SOURCE_ENTITY='FLOCK',
            SOURCE_FILENAME='new_flocks.csv',SOURCE_ROW_NUMBER=n+1,SOURCE_ID=g,TARGET_ID=g,
            VALIDATION_STATUS='READY',RAW_DATA=raw,NORMALIZED_DATA=normalized))
        r._migration_id_map.append(dict(MIGRATION_BATCH_ID=BATCH,MIGRATION_RAW_ROW_ID=f'raw-{n}',SOURCE_ENTITY='FLOCK',SOURCE_ID=g,TARGET_ID=g,STATUS='COMMITTED'))
    r._flocks=pd.DataFrame(rows)
    return r


def preview(r,holds=None):return r.preview_flock_date_repair(BATCH,BASIS,{} if holds is None else holds)
def commit(r,p=None):return r.commit_flock_date_repair(BATCH,BASIS,{},(p or preview(r))['review_token'],confirmed=True,actor=ACTOR)


def test_contracts_preserve_original_and_correct_new_aliases():
    old=load_mapping_for_version('EFNS-EIMS-DATAVERSE-3')
    new=load_mapping()
    assert old['entities']['FLOCK']['field_map']['EST_DISPOSAL']=='new_disposaldate'
    assert new['schema_version']!=old['schema_version']
    assert new['entities']['FLOCK']['field_map']['EST_DISPOSAL']=='new_disposaldate2'
    for contract in (old,new):
        n,_=normalize_migration_row('FLOCK',{'new_disposaldate':'2025-04-01','new_disposaldate2':None},2,contract)
        assert (n['EST_DISPOSAL'],n['DISPOSAL_DATE'])==((date(2025,4,1),None) if contract==old else (None,date(2025,4,1)))
    aliases=_aliases('FLOCK',_mapping())
    assert aliases['NEW_DISPOSALDATE']=='DISPOSAL_DATE'
    assert aliases['NEW_DISPOSALDATE2']=='EST_DISPOSAL'
    assert aliases['ESTIMATED_DISPOSAL_DATE']=='EST_DISPOSAL'
    assert aliases['DISPOSAL_DATE']=='DISPOSAL_DATE'
    assert aliases['EST_DISPOSAL']=='EST_DISPOSAL'


def test_future_full_package_auto_detection_selects_revision_not_legacy_v3():
    from tests.test_eims_migration import _v3_sources
    from data.migration import analyze_migration_files
    sources,old,_=_v3_sources()
    assert old['schema_version']=='EFNS-EIMS-DATAVERSE-3'
    assert analyze_migration_files(sources).schema_version==load_mapping()['schema_version']
    assert analyze_migration_files(sources,old).schema_version==old['schema_version']


def test_preview_only_and_commit_preserve_evidence_and_flock(repo):
    before=repo._flocks.copy(deep=True)
    evidence=deepcopy((repo._migration_batches,repo._migration_raw_rows,repo._migration_id_map))
    p=preview(repo);assert p['counts']=={'READY':2,'NO_CHANGE':2}
    pd.testing.assert_frame_equal(before,repo._flocks)
    result=commit(repo,p);assert result['changed']==2
    assert repo._flocks.iloc[0].EST_DISPOSAL==date(2025,6,1)
    assert repo._flocks.iloc[0].DISPOSAL_DATE==date(2025,5,1)
    assert repo._flocks.iloc[1].DISPOSAL_DATE is None
    assert len(repo._continuity_rows['RECORD_REVISION'])==2
    assert all(r['CHANGE_KIND']==KIND and r['BUSINESS_AT'] is None for r in repo._continuity_rows['RECORD_REVISION'])
    assert len(repo._migration_review_audit)==1
    assert (repo._migration_batches,repo._migration_raw_rows,repo._migration_id_map)==evidence
    for field in ('BIRD_COUNT','STATUS','PLACEMENT_DATE','HATCH_DATE','FACILITY_ID','ACCOUNT_ID'):
        assert repo._flocks[field].tolist()==before[field].tolist()
    assert repo._continuity_rows['FLOCK_EVENT']==[]
    assert commit(repo,p)['changed']==0
    assert len(repo._continuity_rows['RECORD_REVISION'])==2


@pytest.mark.parametrize('kind',['dates','version','revision','event','audit'])
def test_later_changes_are_held(repo,kind):
    g=repo._flocks.iloc[0].FLOCK_ID
    if kind=='dates':repo._flocks.at[0,'DISPOSAL_DATE']=date(2025,12,1)
    if kind=='version':repo._flocks.at[0,'UPDATED_AT']=datetime(2025,10,2)
    if kind=='revision':repo._continuity_rows['RECORD_REVISION'].append(dict(REVISION_ID='later',ENTITY_TYPE='FLOCK',ENTITY_ID=g))
    if kind=='event':repo._continuity_rows['FLOCK_EVENT'].append(dict(EVENT_ID='later',FLOCK_ID=g))
    if kind=='audit':repo._migration_review_audit=[dict(AUDIT_ID='later',ENTITY_ID=g)]
    assert preview(repo)['rows'][0]['status']=='HOLD'


@pytest.mark.parametrize('change',['core','raw','mapping','event','scope','basis'])
def test_stale_preview_fails(repo,change):
    p=preview(repo);basis=BASIS;holds={}
    if change=='core':repo._flocks.at[0,'UPDATED_AT']=datetime(2025,10,3)
    if change=='raw':repo._migration_raw_rows[0]['RAW_DATA']['NEW_DISPOSALDATE']='2025-04-30'
    if change=='mapping':repo._migration_id_map[0]['TARGET_ID']=str(UUID(int=800))
    if change=='event':repo._continuity_rows['FLOCK_EVENT'].append(dict(EVENT_ID='later',FLOCK_ID=repo._flocks.iloc[0].FLOCK_ID))
    if change=='scope':holds={repo._flocks.iloc[0].FLOCK_ID:'Synthetic hold'}
    if change=='basis':basis+=' Changed'
    with pytest.raises(ConcurrencyError):repo.commit_flock_date_repair(BATCH,basis,holds,p['review_token'],confirmed=True,actor=ACTOR)
    assert not repo._continuity_rows['RECORD_REVISION']


def test_audit_failure_rolls_back(repo,monkeypatch):
    before=repo._flocks.copy(deep=True)
    def fail(*args):raise RuntimeError('synthetic mandatory audit failure')
    monkeypatch.setattr(repo,'_date_finish',fail)
    with pytest.raises(RuntimeError):commit(repo)
    pd.testing.assert_frame_equal(before,repo._flocks)
    assert not repo._continuity_rows['RECORD_REVISION']


def test_cas_failure_rolls_back_earlier_rows(repo,monkeypatch):
    before=repo._flocks.copy(deep=True);original=repo._date_update;calls=[]
    def fail_second(tx,old,new):
        calls.append(old)
        if len(calls)==2:raise ConcurrencyError('synthetic competing writer')
        return original(tx,old,new)
    monkeypatch.setattr(repo,'_date_update',fail_second)
    with pytest.raises(ConcurrencyError):commit(repo)
    pd.testing.assert_frame_equal(before,repo._flocks)
    assert not repo._continuity_rows['RECORD_REVISION']


def test_concurrent_same_request_commits_once(repo):
    p=preview(repo)
    with ThreadPoolExecutor(2) as pool:results=list(pool.map(lambda _:commit(repo,p),range(2)))
    assert sum(r['changed'] for r in results)==2
    assert len(repo._migration_review_audit)==1


def test_protected_reversal_and_replay(repo):
    before=repo._flocks.copy(deep=True);result=commit(repo)
    p=repo.preview_flock_date_reversal(result['operation_id'],'Synthetic rollback basis')
    rev=repo.commit_flock_date_reversal(result['operation_id'],p['basis'],p['review_token'],confirmed=True,actor=ACTOR)
    assert rev['changed']==2
    assert repo._flocks['EST_DISPOSAL'].tolist()==before['EST_DISPOSAL'].tolist()
    assert repo._flocks['DISPOSAL_DATE'].tolist()==before['DISPOSAL_DATE'].tolist()
    assert [r['CHANGE_KIND'] for r in repo._continuity_rows['RECORD_REVISION']]==[KIND,KIND,UNDO,UNDO]
    assert repo.commit_flock_date_reversal(result['operation_id'],p['basis'],p['review_token'],confirmed=True,actor=ACTOR)['changed']==0


def test_reversal_refuses_subsequent_changes(repo):
    result=commit(repo);repo._flocks.at[0,'UPDATED_AT']=datetime(2027,1,1)
    with pytest.raises(ConcurrencyError):repo.preview_flock_date_reversal(result['operation_id'],'Synthetic rollback')


def test_admin_authorization_and_actor_injection(repo):
    reader=SimpleNamespace(role=Role.REPORTING_VIEWER,user_id='reader',email='reader@nsegg.ca',active=True)
    with pytest.raises(RepositoryError):AuthorizedRepository(repo,SimpleNamespace(),reader).preview_flock_date_repair(BATCH,BASIS,{})
    admin=SimpleNamespace(role=Role.ADMIN,user_id=ACTOR['user_id'],email=ACTOR['email'],active=True)
    proxy=AuthorizedRepository(repo,SimpleNamespace(),admin)
    p=proxy.preview_flock_date_repair(BATCH,BASIS,{})
    proxy.commit_flock_date_repair(BATCH,BASIS,{},p['review_token'],confirmed=True,actor={'user_id':'forged','email':'forged@nsegg.ca'})
    assert repo._migration_review_audit[0]['ACTOR_EMAIL']==ACTOR['email']


def test_exclusions_and_invalid_identity(repo):
    g=repo._flocks.iloc[0].FLOCK_ID
    assert preview(repo,{g:'Later source modification'})['rows'][0]['status']=='HOLD'
    with pytest.raises(RepositoryError):repo.preview_flock_date_repair(BATCH,BASIS,None)
    with pytest.raises(RepositoryError):preview(repo,{'wrong':'Unknown source'})
    repo._migration_raw_rows.append(deepcopy(repo._migration_raw_rows[0]))
    with pytest.raises(RepositoryError):preview(repo)


def test_old_correction_tokens_invalid_even_if_values_same():
    from data.historical_corrections import HistoricalCorrectionPreview
    p=HistoricalCorrectionPreview('batch','round',[],mapping_revision='old')
    q=HistoricalCorrectionPreview('batch','round',[],mapping_revision='new')
    assert p.review_token!=q.review_token
    assert p.review_token!=__import__('hashlib').sha256(b'[]').hexdigest()


def test_report_and_exports_use_corrected_fields(repo):
    commit(repo)
    frame=build_custom_report(repo,user=user(),dataset='Flocks',selected_fields=['FLOCK_ID','EST_DISPOSAL','DISPOSAL_DATE'])
    # The report reads current CORE values; no export-time column swap.
    row=frame[frame['Flock ID']==repo._flocks.iloc[0].FLOCK_ID].iloc[0]
    assert pd.Timestamp(row['Estimated Disposal Date']).date()==date(2025,6,1)
    csv_data=pd.read_csv(BytesIO(report_csv(user(),frame)),keep_default_na=False)
    wb=load_workbook(BytesIO(report_xlsx(user(),frame)),data_only=True)
    assert len(csv_data)==4 and wb.active.max_row==5
    filtered=build_custom_report(repo,user=user(),dataset='Flocks',selected_fields=['FLOCK_ID','EST_DISPOSAL','DISPOSAL_DATE'],
        filters={'date_from':date(2025,5,1),'date_to':date(2025,5,31),'date_field':'DISPOSAL_DATE'})
    assert len(filtered)==1 and filtered.iloc[0]['Flock ID']==repo._flocks.iloc[0].FLOCK_ID
    with pytest.raises(ValueError,match='date field'):
        build_custom_report(repo,user=user(),dataset='Flocks',selected_fields=['FLOCK_ID'],filters={'date_from':date(2025,5,1)})


def test_rejected_correction_uses_explicit_revision_and_rejects_old_token():
    from tests.test_historical_corrections import _repository,BATCH_ID,_raw
    from data.historical_corrections import build_historical_correction_preview
    r=_repository()
    row=_raw('FLOCK','synthetic-rejected-date',2,dict(FLOCK_ID=str(UUID(int=700)),FLOCK_NUMBER='Synthetic',
        ACCOUNT_ID='account-a',FACILITY_ID='facility-a',HATCH_DATE='2025-01-01',PERMIT_NUMBER='P',BIRD_COUNT='10'))
    row['RAW_DATA'].update(NEW_DISPOSALDATE='2025-05-01',NEW_DISPOSALDATE2='2025-06-01')
    r._migration_raw_rows=[row]
    # _repository's original batch uses a synthetic contract. Supply an explicit
    # v3 contract to prove old RAW's interpretation is revised, not rewritten.
    p=build_historical_correction_preview(BATCH_ID,[row],{'ACCOUNT':{'account-a'},'FACILITY':{'facility-a'}},
        mapping=load_mapping(V3_MAPPING_PATH),facility_accounts={'facility-a':'account-a'})
    assert p.rows[0]['CORRECTED_DATA']['EST_DISPOSAL']==date(2025,6,1)
    assert p.rows[0]['CORRECTED_DATA']['DISPOSAL_DATE']==date(2025,5,1)
    assert p.mapping_revision==load_mapping()['schema_version']
    current=r.preview_historical_correction_review(BATCH_ID)
    old_token=__import__('hashlib').sha256(json.dumps(current.rows,default=str,sort_keys=True).encode()).hexdigest()
    with pytest.raises(ConcurrencyError):r.commit_historical_correction_review(BATCH_ID,old_token,actor=ACTOR)


def test_form_assignment_and_original_workbook_contract_not_swapped():
    from pathlib import Path
    source=(Path(__file__).resolve().parents[1]/'app/pages/5_Flocks.py').read_text()
    assert '"EST_DISPOSAL": estimated_disposal' in source
    assert '"DISPOSAL_DATE": disposal_date' in source
    old=(Path(__file__).resolve().parents[1]/'data/flock_quota_import.py').read_text()
    assert '("ESTIMATED_DISPOSAL_DATE", "EST_DISPOSAL", estimated_disposal)' in old


def test_reversal_blocks_later_business_event(repo):
    operation=commit(repo)['operation_id']
    repo._continuity_rows['FLOCK_EVENT'].append(dict(EVENT_ID='later-event',FLOCK_ID=repo._flocks.iloc[0].FLOCK_ID))
    with pytest.raises(RepositoryError,match='business events'):repo.preview_flock_date_reversal(operation,'Synthetic rollback')


def test_missing_source_column_not_interpreted_as_null(repo):
    del repo._migration_raw_rows[0]['RAW_DATA']['NEW_DISPOSALDATE2']
    assert preview(repo)['rows'][0]['status']=='HOLD'
    assert 'missing' in preview(repo)['rows'][0]['reason']
