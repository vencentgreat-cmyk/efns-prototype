"""Unknown quantities are not zero; fixtures contain no real business data."""

from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.services.report_builder import build_report
from app.services.report_center import build_custom_report, report_csv, report_xlsx
from data.repositories.mock import MockRepository
from tests.auth_support import authenticated_app
from tests.test_report_center import user


@pytest.fixture
def nullable_repo():
    repo = MockRepository()
    groups = ["unknown", "mixed", "zero", "negative", "cancel"]
    repo._accounts = pd.DataFrame([
        {"ACCOUNT_ID": group, "ORGANIZATION_NAME": f"Synthetic {group}"}
        for group in groups
    ])
    records = []
    for group, values in [
        ("unknown", [None, None]), ("mixed", [None, 7]),
        ("zero", [None, 0]), ("negative", [None, -3]),
        ("cancel", [-4, 4]), (None, [None, None]),
    ]:
        for value in values:
            records.append({
                "PRODUCTION_ID": f"synthetic-{len(records)}",
                "PRODUCER_ACCOUNT_ID": group, "FLOCK_ID": None,
                "FACILITY_ID": None, "IMPORT_ID": None,
                "REPORTING_YEAR": 2026, "REPORTING_WEEK": 1,
                "TOTAL": value, "TOTAL_RECEIVED": 100,
            })
    repo._production = pd.DataFrame(records)
    return repo


def grouped(repo, **kwargs):
    return build_custom_report(
        repo, user=user(), dataset="Production Records",
        selected_fields=["PRODUCER_ACCOUNT_ID", "PRODUCER_ACCOUNT_NAME", "TOTAL"],
        group_by=["PRODUCER_ACCOUNT_ID", "PRODUCER_ACCOUNT_NAME"],
        aggregations={"TOTAL": "sum"}, **kwargs,
    )


def read_xlsx(data):
    # pandas trims an entirely blank final row on read; inspect the actual
    # worksheet dimensions so an unassigned all-null group is still checked.
    sheet = load_workbook(BytesIO(data), data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    return pd.DataFrame(rows[1:], columns=rows[0])


def assert_nullable_groups(result):
    assert len(result) == 7
    indexed = result.set_index("Producer Account ID")
    assert pd.isna(indexed.loc["unknown", "Source Total"])
    for group, expected in [("mixed", 7), ("zero", 0), ("negative", -3), ("cancel", 0)]:
        assert indexed.loc[group, "Source Total"] == expected
    unassigned = result[result["Producer Account ID"].isna()]
    assert len(unassigned) == 2 and unassigned["Source Total"].isna().all()
    assert unassigned["Producer identity key"].nunique() == 2


@pytest.mark.parametrize("dtype", ["float64", "Float64", "decimal"])
def test_sum_preserves_null_zero_and_signed_values(nullable_repo, dtype):
    if dtype == "decimal":
        nullable_repo._production["TOTAL"] = nullable_repo._production.TOTAL.map(
            lambda v: None if pd.isna(v) else Decimal(str(v))
        )
    else:
        nullable_repo._production["TOTAL"] = nullable_repo._production.TOTAL.astype(dtype)
    original = nullable_repo._production.copy(deep=True)
    result = grouped(nullable_repo)
    assert_nullable_groups(result)
    pd.testing.assert_frame_equal(original, nullable_repo._production)


def test_csv_xlsx_keep_unknown_blank_and_real_zero_numeric(nullable_repo):
    result = grouped(nullable_repo)
    for exported in [
        pd.read_csv(BytesIO(report_csv(user(), result))),
        read_xlsx(report_xlsx(user(), result)),
    ]:
        assert_nullable_groups(exported)
        pd.testing.assert_frame_equal(
            result.fillna(""), exported.fillna(""), check_dtype=False,
        )


def test_old_builder_retains_unknown_source_total_without_substitution(nullable_repo):
    result = build_report(nullable_repo, ["PRODUCER_ACCOUNT_ID", "SOURCE_TOTAL", "TOTAL_RECEIVED"])
    assert len(result) == 12
    assert result.loc[result["Producer Account ID"] == "unknown", "Source Total"].isna().all()
    assert result.loc[result["Producer Account ID"].isna(), "Source Total"].isna().all()
    assert result["Total Received"].eq(100).all()
    for exported in [pd.read_csv(BytesIO(report_csv(user(), result))), read_xlsx(report_xlsx(user(), result))]:
        pd.testing.assert_frame_equal(result.fillna(""), exported.fillna(""), check_dtype=False)


@pytest.mark.parametrize("operation, expected", [("min", -4), ("max", 4), ("mean", 0)])
def test_other_aggregations_and_filters_unchanged(nullable_repo, operation, expected):
    result = build_custom_report(
        nullable_repo, user=user(), dataset="Production Records",
        selected_fields=["PRODUCER_ACCOUNT_ID", "TOTAL"],
        filters={"account_id": "cancel", "reporting_year": 2026, "reporting_week": 1},
        group_by=["PRODUCER_ACCOUNT_ID"], aggregations={"TOTAL": operation},
    )
    assert len(result) == 1 and result.iloc[0]["Source Total"] == expected


def test_page_preview_and_export_use_nullable_group_result(nullable_repo, tmp_path, monkeypatch):
    app = authenticated_app(Path(__file__).parents[1]/"app/Home.py", tmp_path, monkeypatch)
    app.session_state._base_repo = nullable_repo
    app.session_state.repo = nullable_repo
    app.switch_page("pages/4_Reports.py").run()
    app.radio(key="report_center_section").set_value("Customize Report").run()
    app.selectbox(key="custom_report_dataset").set_value("Production Records").run()
    app.multiselect(key="custom_report_columns").set_value(["Producer Account ID", "Producer Account", "Source Total"]).run()
    app.multiselect(key="custom_report_group").set_value(["Producer Account ID", "Producer Account"]).run()
    app.selectbox(key="custom_report_aggregate_field").set_value("Source Total").run()
    app.selectbox(key="custom_report_aggregate").set_value("sum")
    next(button for button in app.button if button.label == "Run Report").click().run()
    assert not app.exception and not app.error
    result = app.session_state.custom_report_result
    assert_nullable_groups(result)
    assert any(isinstance(table.value, pd.DataFrame) and table.value.equals(result) for table in app.dataframe)
    for exported in [pd.read_csv(BytesIO(report_csv(user(), result))), read_xlsx(report_xlsx(user(), result))]:
        assert_nullable_groups(exported)
