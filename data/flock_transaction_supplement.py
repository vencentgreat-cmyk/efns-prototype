"""Restricted Dataverse associated-view supplement; never performs I/O writes."""
from collections import Counter
from dataclasses import dataclass
import datetime as dt
from decimal import Decimal, InvalidOperation
import hashlib
import math
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qsl
from uuid import UUID
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

from data.source_evidence import evidence_json


ROUTE = "HISTORICAL_FLOCK_TX_SUPPLEMENT"
BUSINESS_FIELDS = ("FLOCK_ID", "TRANSACTION_TYPE", "TRANSACTION_DATE", "QUANTITY",
                   "NOTES", "COUNT_SOURCE_LABEL")
REQUIRED_LOGICAL = {
    "new_flocktransactionid", "new_flockidid", "new_date", "new_birdcount",
    "new_flocktransactiontype", "efc_flockcounttype", "modifiedon", "createdon",
}
NS = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


class SupplementError(ValueError):
    pass


def guid(value):
    text = str(value or "").strip()
    if len(text) != 36:
        raise SupplementError("A canonical source GUID is required.")
    try:
        return str(UUID(text)).upper()
    except ValueError as exc:
        raise SupplementError("A canonical source GUID is required.") from exc


def optional_text(value):
    if isinstance(value, float) and math.isnan(value):
        return None
    return None if value is None or str(value).strip() == "" else str(value).strip()


def quantity(value):
    try:
        if isinstance(value, bool):
            raise InvalidOperation
        result = Decimal(str(value))
        if not result.is_finite() or result < 0 or result > Decimal("9999999999.99"):
            raise InvalidOperation
        if result != result.quantize(Decimal("0.01")):
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise SupplementError("Quantity must be a nonnegative NUMBER(12,2) value.") from exc


def business_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    try:
        return dt.date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise SupplementError("Transaction Date must be an Excel date or ISO date.") from exc


def signature(record):
    """Compare business content, never ingestion timestamps or file hashes."""
    return (guid(record.get("FLOCK_ID")), optional_text(record.get("TRANSACTION_TYPE")),
            business_date(record.get("TRANSACTION_DATE")), quantity(record.get("QUANTITY")).quantize(Decimal('0.01')),
            optional_text(record.get("NOTES")), optional_text(record.get("COUNT_SOURCE_LABEL")))


@dataclass
class ParsedSupplement:
    filename: str
    file_hash: str
    worksheet: str
    rows: list
    metadata: dict


def parse_associated_view(filename, content):
    if Path(filename).suffix.lower() != ".xlsx" or len(content) > 10 * 1024 * 1024:
        raise SupplementError("Use the original associated-view XLSX (maximum 10 MiB).")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise SupplementError("The workbook expands beyond the supported size.")
            for name in archive.namelist():
                if name.endswith(".rels"):
                    if any(node.get("TargetMode") == "External" for node in ET.fromstring(archive.read(name))):
                        raise SupplementError("External workbook relationships are not supported.")
            workbook = load_workbook(BytesIO(content), data_only=False, keep_links=False)
            visible = [sheet for sheet in workbook if sheet.sheet_state == "visible"]
            hidden = [sheet for sheet in workbook if sheet.sheet_state != "visible"]
            if len(visible) != 1 or len(hidden) != 1:
                raise SupplementError("Expected one data worksheet and one hidden metadata worksheet.")
            sheet = visible[0]
            metadata_text = str(hidden[0]["A1"].value or "")
            parts = metadata_text.split(":", 2)
            if len(parts) != 3 or parts[0] != "new_flocktransaction":
                raise SupplementError("Not a Dataverse Flock Transaction associated view.")
            mapping_pairs = parse_qsl(parts[2], keep_blank_values=True)
            mapping = dict(mapping_pairs)
            if len(mapping) != len(mapping_pairs) or not REQUIRED_LOGICAL.issubset(mapping):
                raise SupplementError("Missing or duplicate required logical field metadata.")
            headers = [cell.value for cell in sheet[1]]
            if len(set(headers)) != len(headers):
                raise SupplementError("Duplicate worksheet headings.")
            columns = {logical: headers.index(label) + 1 for logical, label in mapping.items()
                       if logical in REQUIRED_LOGICAL | {"new_comments", "checksumLogicalName", "createdby"}
                       and label in headers}
            if not REQUIRED_LOGICAL.issubset(columns):
                raise SupplementError("A required metadata column is absent from the data worksheet.")
            # Keep actual XML cell payloads (including original numeric date serials).
            xml_cells = {}
            for name in archive.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    tree = ET.fromstring(archive.read(name))
                    first = tree.find(".//x:c[@r='A1']", NS)
                    if first is not None and "".join(first.itertext()) == str(headers[0]):
                        xml_cells = {cell.get("r"): ET.tostring(cell, encoding="unicode")
                                     for cell in tree.findall(".//x:c", NS)}
            if not xml_cells:
                raise SupplementError("Cannot locate the source worksheet XML.")
            rows = []
            for index in range(2, sheet.max_row + 1):
                cells = list(sheet[index])
                if not any(cell.value is not None for cell in cells):
                    continue
                raw = {cell.coordinate: {"value": cell.value, "type": cell.data_type,
                       "number_format": cell.number_format, "xml": xml_cells.get(cell.coordinate)}
                       for cell in cells}
                values = {key: sheet.cell(index, column).value for key, column in columns.items()}
                errors = []
                record = None
                try:
                    if any(cell.data_type == "f" for cell in cells):
                        raise SupplementError("Source formulas are not accepted or evaluated.")
                    record = dict(FLOCK_TRANSACTION_ID=guid(values.get("new_flocktransactionid")),
                        TRANSACTION_TYPE=optional_text(values.get("new_flocktransactiontype")),
                        QUANTITY=quantity(values.get("new_birdcount")),
                        TRANSACTION_DATE=business_date(values.get("new_date")),
                        NOTES=optional_text(values.get("new_comments")),
                        COUNT_SOURCE_LABEL=optional_text(values.get("efc_flockcounttype")))
                    if record["TRANSACTION_TYPE"] not in {"Count", "Delivery"}:
                        raise SupplementError("This supplement supports only Count and Delivery rows.")
                    if record["TRANSACTION_TYPE"] == "Count" and not record["COUNT_SOURCE_LABEL"]:
                        raise SupplementError("Count source label is required in this associated view.")
                    if len(record["NOTES"] or "") > 2000 or len(record["COUNT_SOURCE_LABEL"] or "") > 255:
                        raise SupplementError("Source text exceeds the supported column length.")
                except SupplementError as exc:
                    errors.append(str(exc))
                rows.append(dict(source_row=index, raw=raw, logical_values=values,
                                 record=record, errors=errors))
            counts = Counter(row["record"]["FLOCK_TRANSACTION_ID"] for row in rows if row["record"])
            for row in rows:
                if row["record"] and counts[row["record"]["FLOCK_TRANSACTION_ID"]] > 1:
                    row["errors"].append("Duplicate source transaction GUID in this workbook.")
            if len(rows) != 13:
                raise SupplementError("This restricted supplement requires exactly 13 source rows.")
            return ParsedSupplement(Path(filename).name, hashlib.sha256(content).hexdigest(),
                                    sheet.title, rows, dict(logical_mapping=mapping,
                                    hidden_metadata=metadata_text, excel_epoch=workbook.epoch))
    except SupplementError:
        raise
    except Exception as exc:
        raise SupplementError("Cannot parse the original associated-view workbook.") from exc


def classify(parsed, binding, target, existing, original_batch, account, facility):
    required = {"original_batch_id", "flock_id", "anchor_transaction_id", "basis"}
    if set(binding) != required or not str(binding.get("basis") or "").strip():
        raise SupplementError("Explicit Flock binding, original batch and review basis are required.")
    if len(str(binding['basis'])) > 500:
        raise SupplementError("Keep the explicit binding basis within 500 characters.")
    flock_id, anchor_id = guid(binding["flock_id"]), guid(binding["anchor_transaction_id"])
    if not original_batch or str(original_batch.get("STATUS")).upper() != "COMMITTED":
        raise SupplementError("The explicitly linked original migration batch must be committed.")
    if not target or guid(target.get("FLOCK_ID")) != flock_id or not account or not facility:
        raise SupplementError("The bound Flock, Account or Facility is missing or ambiguous.")
    permits = {optional_text(row["logical_values"].get("new_flockidid")) for row in parsed.rows}
    if permits != {optional_text(target.get("PERMIT_NUMBER"))} or None in permits:
        raise SupplementError("Source lookup labels contradict the explicit Flock binding.")
    indexed = {}
    for item in existing:
        indexed.setdefault(guid(item.get("FLOCK_TRANSACTION_ID")), []).append(item)
    anchor_rows = [row for row in parsed.rows if row["record"] and row["record"]["FLOCK_TRANSACTION_ID"] == anchor_id]
    anchor_existing = indexed.get(anchor_id, [])
    if (len(anchor_rows) != 1 or anchor_rows[0]["errors"] or len(anchor_existing) != 1
            or anchor_rows[0]["record"]["TRANSACTION_TYPE"] != "Delivery"):
        raise SupplementError("One existing same-GUID Delivery anchor is required.")
    rows = []
    for row in parsed.rows:
        record = dict(row["record"], FLOCK_ID=flock_id) if row["record"] else None
        errors = list(row["errors"])
        matches = indexed.get(record["FLOCK_TRANSACTION_ID"], []) if record else []
        status = "INVALID" if errors else "NEW"
        if not errors and matches:
            try:
                same = len(matches) == 1 and signature(matches[0]) == signature(record)
            except (SupplementError, TypeError):
                same = False
            status = "EXISTING_MATCH" if same else "CONFLICT"
        rows.append(dict(row, record=record, classification=status))
    if Counter(row["record"]["TRANSACTION_TYPE"] for row in rows if row["record"]) != {"Count":12,"Delivery":1}:
        # Malformed rows stay visible as INVALID; never drop them to reach 13.
        for row in rows:
            if row["classification"] == "NEW":
                row["classification"] = "INVALID"
                row["errors"].append("Expected 12 Count rows and one Delivery anchor.")
    review_token = hashlib.sha256(evidence_json(dict(file_hash=parsed.file_hash,
        binding=binding, target=target, account=account, facility=facility,
        anchor=signature(anchor_existing[0]), original_batch_id=original_batch["MIGRATION_BATCH_ID"])).encode()).hexdigest()
    counts = {name:sum(row["classification"] == name for row in rows)
              for name in ("NEW","EXISTING_MATCH","CONFLICT","INVALID")}
    return dict(file_hash=parsed.file_hash, worksheet=parsed.worksheet, source_rows=len(rows),
        counts=counts, rows=rows, target=target, account=account, facility=facility,
        binding=dict(binding), review_token=review_token,
        can_commit=not counts["CONFLICT"] and not counts["INVALID"])
