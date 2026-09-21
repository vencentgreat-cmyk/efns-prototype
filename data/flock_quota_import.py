"""In-memory parsing and validation for routine Flock and Quota imports."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
import datetime as dt
import hashlib
import io
import math
from typing import Any
from uuid import uuid4
from zipfile import BadZipFile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from data.constants import (
    EGG_COLOURS,
    FLOCK_QUOTA_TYPES,
    FLOCK_STATUSES,
    QUOTA_TYPES,
    UNASSIGNED_LABEL,
)


WORKSHEET_NAME = "Batch Import"
MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_BATCH_ROWS = 5_000
SUPPORTED_RECORD_TYPES = ("QUOTA_REGISTRATION", "FLOCK")
SUPPORTED_ACTIONS = ("CREATE", "UPDATE")
DEV_IDENTIFIER_PREFIX = "DEV_DEMO_"

EXPECTED_HEADERS = (
    "BATCH_ID",
    "RECORD_TYPE",
    "ACTION",
    "ACCOUNT",
    "FACILITY",
    "FACILITY_DETAIL",
    "REGISTRATION_NUMBER",
    "QUOTA_NAME",
    "QUOTA_TYPE",
    "QUOTA_REGISTRATION",
    "FLOCK_QUOTA_TYPE",
    "FLOCK_NUMBER",
    "FLOCK_STATUS",
    "CREATE_DELIVERY_TRANSACTION",
    "PERMIT_NUMBER",
    "HATCH_DATE",
    "BIRD_COUNT",
    "EGG_COLOUR",
    "PERMIT_DATE",
    "DATE_ORDERED",
    "PLACEMENT_DATE",
    "ESTIMATED_DISPOSAL_DATE",
    "DISPOSAL_DATE",
    "BIRDS_DISPOSED",
    "BIRD_STRAIN",
    "EFFECTIVE_DATE",
    "END_DATE",
    "SOURCE_REFERENCE",
    "COMMENTS",
)


class WorkbookImportError(ValueError):
    """Raised when an uploaded workbook cannot be safely parsed."""


def _missing(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _text(value: Any) -> str | None:
    if _missing(value):
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _key(value: Any) -> str:
    return (_text(value) or "").casefold()


def _date_value(value: Any, label: str, errors: list[str]) -> dt.date | None:
    if _missing(value):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    converted = pd.to_datetime(value, errors="coerce")
    if pd.isna(converted):
        errors.append(f"{label} is not a valid date.")
        return None
    return converted.date()


def _integer_value(value: Any, label: str, errors: list[str]) -> int | None:
    if _missing(value):
        return None
    if isinstance(value, bool):
        errors.append(f"{label} must be an integer.")
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        errors.append(f"{label} must be an integer.")
        return None
    if not math.isfinite(number) or not number.is_integer():
        errors.append(f"{label} must be an integer.")
        return None
    return int(number)


def _boolean_value(value: Any, label: str, errors: list[str]) -> bool | None:
    if _missing(value):
        return None
    if isinstance(value, bool):
        return value
    normalized = _key(value)
    if normalized in {"true", "yes", "y", "1"}:
        return True
    if normalized in {"false", "no", "n", "0"}:
        return False
    errors.append(f"{label} must be true or false.")
    return None


@dataclass
class SourceRow:
    row_number: int
    values: dict[str, Any]
    blank: bool = False


@dataclass
class ParsedBatch:
    filename: str
    file_hash: str
    file_size: int
    rows: list[SourceRow]
    global_errors: list[str] = field(default_factory=list)
    global_warnings: list[str] = field(default_factory=list)


@dataclass
class ValidatedRow:
    row_number: int
    record_type: str
    action: str
    business_identifier: str | None
    status: str
    messages: list[str]
    resolved_target_identifier: str | None = None
    changes: dict[str, tuple[Any, Any]] = field(default_factory=dict)
    write_record: dict[str, Any] | None = None

    @property
    def message(self) -> str:
        return " ".join(self.messages)

    @property
    def change_summary(self) -> str:
        return "; ".join(
            f"{name}: {old if old is not None else '(blank)'} -> {new if new is not None else '(blank)'}"
            for name, (old, new) in self.changes.items()
        )


@dataclass
class ValidationReport:
    parsed: ParsedBatch
    batch_id: str | None
    rows: list[ValidatedRow]
    global_errors: list[str] = field(default_factory=list)
    global_warnings: list[str] = field(default_factory=list)

    @property
    def rejected_count(self) -> int:
        return sum(row.status == "Rejected" for row in self.rows)

    @property
    def warning_count(self) -> int:
        return sum(row.status == "Warning" for row in self.rows)

    @property
    def ready_count(self) -> int:
        return sum(row.status == "Ready" for row in self.rows)

    @property
    def record_rows(self) -> list[ValidatedRow]:
        return [row for row in self.rows if row.record_type in SUPPORTED_RECORD_TYPES]

    @property
    def can_import(self) -> bool:
        return bool(self.record_rows) and not self.global_errors and self.rejected_count == 0

    @property
    def quota_records(self) -> list[dict[str, Any]]:
        return [
            dict(row.write_record or {})
            for row in self.rows
            if row.record_type == "QUOTA_REGISTRATION" and row.status != "Rejected" and row.write_record
        ]

    @property
    def flock_records(self) -> list[dict[str, Any]]:
        return [
            dict(row.write_record or {})
            for row in self.rows
            if row.record_type == "FLOCK" and row.status != "Rejected" and row.write_record
        ]

    def to_frame(self) -> pd.DataFrame:
        return pd.DataFrame(
            [
                {
                    "ROW_NUMBER": row.row_number,
                    "RECORD_TYPE": row.record_type or "BLANK_ROW",
                    "ACTION": row.action,
                    "BUSINESS_IDENTIFIER": row.business_identifier,
                    "VALIDATION_STATUS": row.status,
                    "MESSAGE": row.message,
                    "RESOLVED_TARGET_IDENTIFIER": row.resolved_target_identifier,
                    "CHANGES": row.change_summary,
                }
                for row in self.rows
            ]
        )


def parse_flock_quota_workbook(file_bytes: bytes, filename: str) -> ParsedBatch:
    """Parse the required worksheet without writing the upload to disk."""
    if not file_bytes:
        raise WorkbookImportError("The uploaded workbook is empty.")
    if len(file_bytes) > MAX_WORKBOOK_BYTES:
        raise WorkbookImportError("The workbook exceeds the 10 MB import limit.")
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except (BadZipFile, EOFError, InvalidFileException, OSError, TypeError, ValueError, KeyError) as exc:
        raise WorkbookImportError(
            "The workbook is unreadable, malformed, or password-protected."
        ) from exc
    if WORKSHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise WorkbookImportError(f"The workbook must contain a '{WORKSHEET_NAME}' worksheet.")

    worksheet = workbook[WORKSHEET_NAME]
    iterator = worksheet.iter_rows(values_only=True)
    try:
        original_headers = next(iterator)
    except StopIteration as exc:
        workbook.close()
        raise WorkbookImportError(f"The '{WORKSHEET_NAME}' worksheet is empty.") from exc

    normalized_headers = [(_text(value) or "").strip().upper() for value in original_headers]
    global_errors: list[str] = []
    global_warnings: list[str] = []
    nonblank_headers = [header for header in normalized_headers if header]
    duplicates = sorted(name for name, count in Counter(nonblank_headers).items() if count > 1)
    missing_headers = sorted(set(EXPECTED_HEADERS) - set(nonblank_headers))
    unexpected_headers = sorted(set(nonblank_headers) - set(EXPECTED_HEADERS))
    whitespace_headers = sorted(
        str(value) for value in original_headers if isinstance(value, str) and value != value.strip()
    )
    if duplicates:
        global_errors.append("Duplicate headers: " + ", ".join(duplicates) + ".")
    if missing_headers:
        global_errors.append("Missing required headers: " + ", ".join(missing_headers) + ".")
    if unexpected_headers:
        global_warnings.append("Unexpected headers: " + ", ".join(unexpected_headers) + ".")
    if whitespace_headers:
        global_warnings.append(
            "Headers with leading or trailing whitespace: " + ", ".join(whitespace_headers) + "."
        )

    rows: list[SourceRow] = []
    for row_number, values in enumerate(iterator, start=2):
        if row_number - 1 > MAX_BATCH_ROWS:
            workbook.close()
            raise WorkbookImportError(f"The workbook exceeds the {MAX_BATCH_ROWS:,}-row import limit.")
        blank = all(_missing(value) for value in values)
        mapped: dict[str, Any] = {}
        for index, header in enumerate(normalized_headers):
            if header and header not in mapped:
                mapped[header] = values[index] if index < len(values) else None
        rows.append(SourceRow(row_number=row_number, values=mapped, blank=blank))
    workbook.close()
    if not rows:
        global_errors.append("The worksheet contains no data rows.")
    return ParsedBatch(
        filename=filename,
        file_hash=hashlib.sha256(file_bytes).hexdigest(),
        file_size=len(file_bytes),
        rows=rows,
        global_errors=global_errors,
        global_warnings=global_warnings,
    )


def _frame_matches(frame: pd.DataFrame, column: str, value: Any) -> pd.DataFrame:
    if frame.empty or column not in frame or _missing(value):
        return frame.iloc[0:0].copy()
    target = _key(value)
    return frame[frame[column].fillna("").astype(str).str.strip().str.casefold() == target]


def _one_match(
    frame: pd.DataFrame,
    id_column: str,
    name_column: str,
    value: Any,
    label: str,
    errors: list[str],
) -> dict[str, Any] | None:
    if _missing(value):
        return None
    by_id = _frame_matches(frame, id_column, value)
    matches = by_id if not by_id.empty else _frame_matches(frame, name_column, value)
    if matches.empty:
        errors.append(f"{label} was not found.")
        return None
    if len(matches) > 1:
        errors.append(f"{label} is ambiguous and must resolve to one record.")
        return None
    return matches.iloc[0].to_dict()


def _changes(current: dict[str, Any], values: dict[str, Any], ignored: set[str]) -> dict[str, tuple[Any, Any]]:
    changed: dict[str, tuple[Any, Any]] = {}
    for name, new_value in values.items():
        if name in ignored:
            continue
        old_value = current.get(name)
        old_missing = _missing(old_value)
        new_missing = _missing(new_value)
        if old_missing and new_missing:
            continue
        if isinstance(old_value, (dt.date, dt.datetime)) or isinstance(new_value, (dt.date, dt.datetime)):
            old_compare = pd.to_datetime(old_value, errors="coerce")
            new_compare = pd.to_datetime(new_value, errors="coerce")
            if not (pd.isna(old_compare) and pd.isna(new_compare)) and old_compare != new_compare:
                changed[name] = (old_value, new_value)
        elif old_value != new_value:
            changed[name] = (old_value, new_value)
    return changed


def _status(errors: list[str], warnings: list[str]) -> str:
    if errors:
        return "Rejected"
    return "Warning" if warnings else "Ready"


def validate_flock_quota_batch(parsed: ParsedBatch, repo) -> ValidationReport:
    """Validate every source row using repository data without performing writes."""
    accounts = repo.get_accounts()
    facilities = repo.get_facilities()
    facility_details = repo.get_facility_details()
    quotas = repo.get_quota_registrations()
    flocks = repo.get_flocks()

    nonblank = [row for row in parsed.rows if not row.blank]
    batch_values = [_text(row.values.get("BATCH_ID")) for row in nonblank]
    distinct_batches = sorted({value for value in batch_values if value})
    global_errors = list(parsed.global_errors)
    global_warnings = list(parsed.global_warnings)
    if any(not value for value in batch_values):
        global_errors.append("BATCH_ID is required on every record row.")
    if len(distinct_batches) > 1:
        global_errors.append("All record rows must use the same BATCH_ID.")
    batch_id = distinct_batches[0] if len(distinct_batches) == 1 else None

    registration_counts = Counter(
        _key(row.values.get("REGISTRATION_NUMBER"))
        for row in nonblank
        if _key(row.values.get("RECORD_TYPE")) == "quota_registration"
        and _key(row.values.get("REGISTRATION_NUMBER"))
    )
    flock_counts = Counter(
        _key(row.values.get("FLOCK_NUMBER"))
        for row in nonblank
        if _key(row.values.get("RECORD_TYPE")) == "flock" and _key(row.values.get("FLOCK_NUMBER"))
    )
    permit_counts = Counter(
        _key(row.values.get("PERMIT_NUMBER"))
        for row in nonblank
        if _key(row.values.get("RECORD_TYPE")) == "flock" and _key(row.values.get("PERMIT_NUMBER"))
    )

    results_by_row: dict[int, ValidatedRow] = {}
    same_batch_quotas: dict[str, dict[str, Any]] = {}

    for source in parsed.rows:
        if source.blank:
            results_by_row[source.row_number] = ValidatedRow(
                source.row_number, "", "", None, "Warning", ["Completely blank row."]
            )
            continue
        record_type = (_text(source.values.get("RECORD_TYPE")) or "").upper()
        action = (_text(source.values.get("ACTION")) or "").upper()
        if record_type != "QUOTA_REGISTRATION":
            continue
        errors: list[str] = []
        warnings: list[str] = []
        if action not in SUPPORTED_ACTIONS:
            errors.append("ACTION must be CREATE or UPDATE.")
        registration = _text(source.values.get("REGISTRATION_NUMBER"))
        if not registration:
            errors.append("Registration Number is required.")
        elif not registration.startswith(DEV_IDENTIFIER_PREFIX):
            errors.append(f"Registration Number must start with {DEV_IDENTIFIER_PREFIX} in this DEV prototype.")
        if registration and registration_counts[_key(registration)] > 1:
            errors.append("Registration Number is duplicated inside this workbook.")
        existing_matches = _frame_matches(quotas, "REGISTRATION_NUMBER", registration)
        current = existing_matches.iloc[0].to_dict() if len(existing_matches) == 1 else None
        if len(existing_matches) > 1:
            errors.append("Registration Number resolves to multiple existing records.")
        elif action == "CREATE" and current:
            errors.append("Registration Number already exists and cannot be created.")
        elif action == "UPDATE" and current is None:
            errors.append("Registration Number does not exist and cannot be updated.")

        account_value = source.values.get("ACCOUNT")
        account = _one_match(accounts, "ACCOUNT_ID", "ORGANIZATION_NAME", account_value, "Account", errors)
        if account is None and current and _missing(account_value):
            account = _one_match(accounts, "ACCOUNT_ID", "ORGANIZATION_NAME", current.get("ACCOUNT_ID"), "Account", errors)
        if account is None and _missing(account_value) and not current:
            errors.append("Account is required.")

        values: dict[str, Any] = {}
        if account:
            values["ACCOUNT_ID"] = account["ACCOUNT_ID"]
        quota_name = _text(source.values.get("QUOTA_NAME"))
        quota_type = _text(source.values.get("QUOTA_TYPE"))
        effective = _date_value(source.values.get("EFFECTIVE_DATE"), "Effective Date", errors)
        end_date = _date_value(source.values.get("END_DATE"), "End Date", errors)
        comments = _text(source.values.get("COMMENTS"))
        if quota_type is not None and quota_type not in QUOTA_TYPES:
            errors.append("Quota Type must use a supported choice.")
        if end_date and effective and end_date < effective:
            errors.append("End Date cannot be earlier than Effective Date.")

        if action == "CREATE":
            values.update(
                {
                    "QUOTA_ID": str(uuid4()),
                    "REGISTRATION_NUMBER": registration,
                    "QUOTA_NAME": quota_name or registration,
                    "QUOTA_TYPE": quota_type,
                    "STATUS": "Active",
                    "EFFECTIVE_DATE": effective,
                    "END_DATE": end_date,
                    "COMMENTS": comments,
                }
            )
            if quota_type is None:
                errors.append("Quota Type is required.")
        elif current:
            values = {"QUOTA_ID": current["QUOTA_ID"], "REGISTRATION_NUMBER": registration}
            if account:
                values["ACCOUNT_ID"] = account["ACCOUNT_ID"]
            for header, field_name, converted in (
                ("QUOTA_NAME", "QUOTA_NAME", quota_name),
                ("QUOTA_TYPE", "QUOTA_TYPE", quota_type),
                ("EFFECTIVE_DATE", "EFFECTIVE_DATE", effective),
                ("END_DATE", "END_DATE", end_date),
                ("COMMENTS", "COMMENTS", comments),
            ):
                if not _missing(source.values.get(header)):
                    values[field_name] = converted
            values["EXPECTED_UPDATED_AT"] = current.get("UPDATED_AT")

        changes = _changes(current or {}, values, {"QUOTA_ID", "EXPECTED_UPDATED_AT"}) if current else {}
        if action == "UPDATE" and current and not changes:
            errors.append("UPDATE does not contain any supported field changes.")
        status = _status(errors, warnings)
        result = ValidatedRow(
            source.row_number,
            record_type,
            action,
            registration,
            status,
            [*errors, *warnings],
            (current or values).get("QUOTA_ID") if current or values else None,
            changes,
            dict(values, ACTION=action) if status != "Rejected" else None,
        )
        results_by_row[source.row_number] = result
        if status != "Rejected" and registration:
            prior = current or {}
            same_batch_quotas[_key(registration)] = {
                "QUOTA_ID": values.get("QUOTA_ID") or prior.get("QUOTA_ID"),
                "ACCOUNT_ID": values.get("ACCOUNT_ID") or prior.get("ACCOUNT_ID"),
                "REGISTRATION_NUMBER": registration,
                "STATUS": values.get("STATUS") or prior.get("STATUS"),
                "EFFECTIVE_DATE": values.get("EFFECTIVE_DATE", prior.get("EFFECTIVE_DATE")),
                "END_DATE": values.get("END_DATE", prior.get("END_DATE")),
            }

    for source in parsed.rows:
        if source.blank or source.row_number in results_by_row:
            continue
        record_type = (_text(source.values.get("RECORD_TYPE")) or "").upper()
        action = (_text(source.values.get("ACTION")) or "").upper()
        errors: list[str] = []
        warnings: list[str] = []
        if record_type not in SUPPORTED_RECORD_TYPES:
            errors.append("RECORD_TYPE must be QUOTA_REGISTRATION or FLOCK.")
        if action not in SUPPORTED_ACTIONS:
            errors.append("ACTION must be CREATE or UPDATE.")
        flock_number = _text(source.values.get("FLOCK_NUMBER"))
        if record_type != "FLOCK":
            results_by_row[source.row_number] = ValidatedRow(
                source.row_number, record_type, action, flock_number, "Rejected", errors
            )
            continue
        if not flock_number:
            errors.append("Flock Number is required.")
        elif not flock_number.startswith(DEV_IDENTIFIER_PREFIX):
            errors.append(f"Flock Number must start with {DEV_IDENTIFIER_PREFIX} in this DEV prototype.")
        if flock_number and flock_counts[_key(flock_number)] > 1:
            errors.append("Flock Number is duplicated inside this workbook.")
        existing_matches = _frame_matches(flocks, "FLOCK_NUMBER", flock_number)
        current = existing_matches.iloc[0].to_dict() if len(existing_matches) == 1 else None
        if len(existing_matches) > 1:
            errors.append("Flock Number resolves to multiple existing records.")
        elif action == "CREATE" and current:
            errors.append("Flock Number already exists and cannot be created.")
        elif action == "UPDATE" and current is None:
            errors.append("Flock Number does not exist and cannot be updated.")

        account_value = source.values.get("ACCOUNT")
        account = _one_match(accounts, "ACCOUNT_ID", "ORGANIZATION_NAME", account_value, "Account", errors)
        if account is None and current and _missing(account_value):
            account = _one_match(accounts, "ACCOUNT_ID", "ORGANIZATION_NAME", current.get("ACCOUNT_ID"), "Account", errors)
        if account is None and _missing(account_value) and not current:
            errors.append("Account is required.")

        facility_value = source.values.get("FACILITY")
        facility = None
        if account:
            scoped_facilities = facilities[facilities["ACCOUNT_ID"] == account["ACCOUNT_ID"]]
            lookup_value = current.get("FACILITY_ID") if current and _missing(facility_value) else facility_value
            facility = _one_match(scoped_facilities, "FACILITY_ID", "FACILITY_NAME", lookup_value, "Facility", errors)
        if facility is None and _missing(facility_value) and not current:
            errors.append("Facility is required.")

        detail_value = source.values.get("FACILITY_DETAIL")
        detail_id = current.get("FACILITY_DETAIL_ID") if current else None
        if not _missing(detail_value):
            if _key(detail_value) == _key(UNASSIGNED_LABEL):
                detail_id = None
            elif facility:
                scoped_details = facility_details[facility_details["FACILITY_ID"] == facility["FACILITY_ID"]]
                detail = _one_match(
                    scoped_details, "FACILITY_DETAIL_ID", "DETAIL_NAME", detail_value, "Facility Detail", errors
                )
                detail_id = detail.get("FACILITY_DETAIL_ID") if detail else None

        quota_value = source.values.get("QUOTA_REGISTRATION")
        quota = None
        if not _missing(quota_value):
            quota = same_batch_quotas.get(_key(quota_value))
            if quota is None:
                quota = _one_match(
                    quotas, "QUOTA_ID", "REGISTRATION_NUMBER", quota_value, "Quota Registration", errors
                )
            if quota and account and quota.get("ACCOUNT_ID") != account.get("ACCOUNT_ID"):
                errors.append("Quota Registration must belong to the selected account.")
            if quota and quota.get("STATUS", "Active") != "Active":
                errors.append("Quota Registration must be active.")
            if quota:
                effective_date = _date_value(quota.get("EFFECTIVE_DATE"), "Quota Effective Date", errors)
                quota_end_date = _date_value(quota.get("END_DATE"), "Quota End Date", errors)
                today = dt.date.today()
                if effective_date and effective_date > today:
                    errors.append("Quota Registration is not yet effective.")
                if quota_end_date and quota_end_date < today:
                    errors.append("Quota Registration has ended.")
        elif current and current.get("QUOTA_ID"):
            matches = _frame_matches(quotas, "QUOTA_ID", current.get("QUOTA_ID"))
            quota = matches.iloc[0].to_dict() if len(matches) == 1 else None

        permit_number = _text(source.values.get("PERMIT_NUMBER"))
        if permit_number is None and current:
            permit_number = _text(current.get("PERMIT_NUMBER"))
        if not permit_number:
            errors.append("Permit Number is required.")
        elif not permit_number.startswith(DEV_IDENTIFIER_PREFIX):
            errors.append(f"Permit Number must start with {DEV_IDENTIFIER_PREFIX} in this DEV prototype.")
        if permit_number and permit_counts[_key(permit_number)] > 1:
            errors.append("Permit Number is duplicated inside this workbook.")
        permit_matches = _frame_matches(flocks, "PERMIT_NUMBER", permit_number)
        if current is not None:
            permit_matches = permit_matches[permit_matches["FLOCK_ID"] != current["FLOCK_ID"]]
        if not permit_matches.empty:
            errors.append("Permit Number is already used by another Flock.")

        status_value = _text(source.values.get("FLOCK_STATUS"))
        quota_type = _text(source.values.get("FLOCK_QUOTA_TYPE"))
        egg_colour = _text(source.values.get("EGG_COLOUR"))
        hatch = _date_value(source.values.get("HATCH_DATE"), "Hatch Date", errors)
        permit_date = _date_value(source.values.get("PERMIT_DATE"), "Permit Date", errors)
        ordered = _date_value(source.values.get("DATE_ORDERED"), "Date Ordered", errors)
        placement = _date_value(source.values.get("PLACEMENT_DATE"), "Placement Date", errors)
        estimated_disposal = _date_value(
            source.values.get("ESTIMATED_DISPOSAL_DATE"), "Estimated Disposal Date", errors
        )
        disposal = _date_value(source.values.get("DISPOSAL_DATE"), "Disposal Date", errors)
        bird_count = _integer_value(source.values.get("BIRD_COUNT"), "Bird Count", errors)
        birds_disposed = _integer_value(source.values.get("BIRDS_DISPOSED"), "Birds Disposed", errors)
        create_delivery = _boolean_value(
            source.values.get("CREATE_DELIVERY_TRANSACTION"), "Create Delivery Transaction", errors
        )
        if status_value is not None and status_value not in FLOCK_STATUSES:
            errors.append("Flock Status must use a supported choice.")
        if quota_type is not None and quota_type not in FLOCK_QUOTA_TYPES:
            errors.append("Flock Quota Type must use a supported choice.")
        if egg_colour is not None and egg_colour not in EGG_COLOURS:
            errors.append("Egg Colour must use a supported choice.")

        merged = dict(current or {})
        values: dict[str, Any] = {
            "FLOCK_ID": current.get("FLOCK_ID") if current else str(uuid4()),
            "FLOCK_NUMBER": flock_number,
        }
        source_map = (
            ("FLOCK_STATUS", "STATUS", status_value),
            ("FLOCK_QUOTA_TYPE", "FLOCK_QUOTA_TYPE", quota_type),
            ("PERMIT_NUMBER", "PERMIT_NUMBER", permit_number),
            ("HATCH_DATE", "HATCH_DATE", hatch),
            ("BIRD_COUNT", "BIRD_COUNT", bird_count),
            ("EGG_COLOUR", "EGG_COLOUR", egg_colour),
            ("PERMIT_DATE", "PERMIT_DATE", permit_date),
            ("DATE_ORDERED", "DATE_ORDERED", ordered),
            ("PLACEMENT_DATE", "PLACEMENT_DATE", placement),
            ("ESTIMATED_DISPOSAL_DATE", "EST_DISPOSAL", estimated_disposal),
            ("DISPOSAL_DATE", "DISPOSAL_DATE", disposal),
            ("BIRDS_DISPOSED", "BIRDS_DISPOSED", birds_disposed),
            ("BIRD_STRAIN", "BIRD_STRAIN", _text(source.values.get("BIRD_STRAIN"))),
            ("COMMENTS", "COMMENTS", _text(source.values.get("COMMENTS"))),
            ("CREATE_DELIVERY_TRANSACTION", "CREATE_DELIVERY_TRANSACTION", create_delivery),
        )
        if account:
            values["ACCOUNT_ID"] = account["ACCOUNT_ID"]
        if facility:
            values["FACILITY_ID"] = facility["FACILITY_ID"]
        if not _missing(detail_value) or action == "CREATE":
            values["FACILITY_DETAIL_ID"] = detail_id
        if not _missing(quota_value) or action == "CREATE":
            values["QUOTA_ID"] = quota.get("QUOTA_ID") if quota else None
            values["QUOTA_REGISTRATION_NUMBER"] = _text(quota_value)
        for header, field_name, converted in source_map:
            if action == "CREATE" or not _missing(source.values.get(header)):
                values[field_name] = converted
        if action == "CREATE":
            values.setdefault("STATUS", "Planned")
            values.setdefault("BIRDS_DISPOSED", 0)
            values.setdefault("CREATE_DELIVERY_TRANSACTION", False)
        elif current:
            values["EXPECTED_UPDATED_AT"] = current.get("UPDATED_AT")
        merged.update({key: value for key, value in values.items() if key not in {"EXPECTED_UPDATED_AT", "QUOTA_REGISTRATION_NUMBER"}})

        if _missing(merged.get("HATCH_DATE")):
            errors.append("Hatch Date is required.")
        if _missing(merged.get("BIRD_COUNT")):
            errors.append("Bird Count is required.")
        final_bird_count = _integer_value(merged.get("BIRD_COUNT"), "Bird Count", errors)
        final_disposed = _integer_value(merged.get("BIRDS_DISPOSED"), "Birds Disposed", errors) or 0
        if final_bird_count is not None and final_bird_count < 0:
            errors.append("Bird Count cannot be negative.")
        if final_disposed < 0:
            errors.append("Birds Disposed cannot be negative.")
        if final_bird_count is not None and final_disposed > final_bird_count:
            errors.append("Birds Disposed cannot exceed Bird Count.")
        final_hatch = _date_value(merged.get("HATCH_DATE"), "Hatch Date", errors)
        final_placement = _date_value(merged.get("PLACEMENT_DATE"), "Placement Date", errors)
        final_disposal = _date_value(merged.get("DISPOSAL_DATE"), "Disposal Date", errors)
        if final_disposal and final_hatch and final_disposal < final_hatch:
            errors.append("Disposal Date cannot be earlier than Hatch Date.")
        if final_disposal and final_placement and final_disposal < final_placement:
            errors.append("Disposal Date cannot be earlier than Placement Date.")

        changes = _changes(
            current or {}, values, {"FLOCK_ID", "EXPECTED_UPDATED_AT", "QUOTA_REGISTRATION_NUMBER"}
        ) if current else {}
        if action == "UPDATE" and current and not changes:
            errors.append("UPDATE does not contain any supported field changes.")
        status = _status(errors, warnings)
        results_by_row[source.row_number] = ValidatedRow(
            source.row_number,
            record_type,
            action,
            flock_number,
            status,
            [*errors, *warnings],
            (current or values).get("FLOCK_ID") if current or values else None,
            changes,
            dict(values, ACTION=action) if status != "Rejected" else None,
        )

    return ValidationReport(
        parsed=parsed,
        batch_id=batch_id,
        rows=[results_by_row[row.row_number] for row in parsed.rows],
        global_errors=global_errors,
        global_warnings=global_warnings,
    )
